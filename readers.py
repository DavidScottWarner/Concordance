"""Turn txt, Markdown, Word, PDF, and HTML files into plain text."""

from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf", ".html", ".htm"}


def read_text_file(path):
    """Read a text file, including one saved with a hidden encoding mark."""
    data = Path(path).read_bytes()
    return decode_bytes(data)


def decode_bytes(data):
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_docx(path):
    try:
        from docx import Document
    except ImportError as error:
        raise RuntimeError("Word documents need the python-docx library.") from error

    document = Document(str(path))
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            chunks.append(" ".join(cell for cell in cells if cell))
    return "\n".join(chunks)


def read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("PDF files need the pypdf library.") from error

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def read_html(path):
    try:
        from bs4 import BeautifulSoup
    except ImportError as error:
        raise RuntimeError("HTML files need the beautifulsoup4 library.") from error

    soup = BeautifulSoup(read_text_file(path), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text("\n")


def read_document(path):
    """Return the plain text of one supported file."""
    extension = Path(path).suffix.lower()
    if extension in {".txt", ".md"}:
        return read_text_file(path)
    if extension == ".docx":
        return read_docx(path)
    if extension == ".pdf":
        return read_pdf(path)
    if extension in {".html", ".htm"}:
        return read_html(path)
    raise RuntimeError(f"This tool cannot read {extension or 'this type of'} files.")


def read_beginning(path, limit=20000):
    """Read enough of a file to see whether it starts with header lines."""
    path = Path(path)
    if path.suffix.lower() in {".txt", ".md"}:
        with path.open("rb") as handle:
            return decode_bytes(handle.read(limit))
    return read_document(path)[:limit]
