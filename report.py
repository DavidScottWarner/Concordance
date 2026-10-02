"""Write one dated results folder for a run."""

import csv
import re
import shutil
from datetime import datetime
from pathlib import Path

from terms import write_terms
from variants import DECISION_COLUMNS, decisions_path, read_decision_rows
from context_words import CONTEXT_COLUMNS, CONTEXT_NOTE, CONTEXT_SKIPPED_NOTE
from corpus import RunStopped, check_stop

RESERVED_COLUMNS = {
    "file_name",
    "file_path",
    "language",
    "short",
    "category",
    "term",
    "location",
    "sentence",
    "keep",
    "notes",
}


def include_in_counts(language, choice):
    choice = (choice or "").strip()
    if choice.casefold() in {"", "all", "all languages"}:
        return True
    return language.casefold() == choice.casefold()


def per_10000(matches, words):
    if words <= 0:
        return ""
    return f"{matches / words * 10000:.2f}"


def output_column_name(name):
    if name in RESERVED_COLUMNS:
        return f"metadata_{name}"
    return name


def metadata_columns(documents):
    columns = []
    for document in documents:
        for key in document.metadata:
            if key not in columns:
                columns.append(key)
    return columns


def group_value(document, column):
    if not column:
        return ""
    value = (document.metadata.get(column) or "").strip()
    return value or "(blank)"


_RUN_STAMP = re.compile(r"^(\d{4}-\d{2}-\d{2})_(\d{2})-(\d{2})(?:-(\d{2}))?_")


def _run_moment(path):
    """When a run folder was created, taken from its name when that name has a time."""
    match = _RUN_STAMP.match(path.name)
    if match:
        second = match.group(4) or "00"
        return datetime.strptime(
            f"{match.group(1)} {match.group(2)}:{match.group(3)}:{second}",
            "%Y-%m-%d %H:%M:%S",
        )
    try:
        return datetime.fromtimestamp(path.stat().st_mtime)
    except OSError:
        return datetime.min


def earlier_runs(results_folder, destination):
    """Finished runs in this results folder, newest first, other than the one being written."""
    folder = Path(results_folder)
    if not folder.is_dir():
        return []
    skip = destination.resolve() if destination is not None else None
    candidates = []
    for path in folder.iterdir():
        if not path.is_dir():
            continue
        try:
            if skip is not None and path.resolve() == skip:
                continue
        except OSError:
            continue
        if (path / "results.xlsx").is_file() or (path / "matches.csv").is_file() or (
            path / "terms_used.csv"
        ).is_file():
            candidates.append(path)
    candidates.sort(key=lambda path: (_run_moment(path), path.name), reverse=True)
    return candidates


def _mark_key(row):
    return (
        _cell_text(row.get("file_path")),
        _cell_text(row.get("category")),
        _cell_text(row.get("term")),
        _cell_text(row.get("location")),
        _cell_text(row.get("sentence")),
    )


def _cell_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _marks_from_rows(rows):
    """Map each match to its keep and notes text. A later row for the same match wins."""
    marks = {}
    for row in rows:
        keep = _cell_text(row.get("keep"))
        notes = _cell_text(row.get("notes"))
        if not keep and not notes:
            continue
        marks[_mark_key(row)] = (keep, notes)
    return marks


def _loose_key(row):
    """Identity shared by the workbook and the CSV when the workbook has no file path."""
    return (
        _cell_text(row.get("file_name")),
        _cell_text(row.get("category")),
        _cell_text(row.get("term")),
        _cell_text(row.get("location")),
        _cell_text(row.get("sentence")),
    )


def _read_csv_rows(path):
    if not path.is_file():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _read_workbook_rows(path):
    """Rows from the Matches sheet. None means the workbook could not be read."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        if "Matches" not in workbook.sheetnames:
            return None
        sheet = workbook["Matches"]
        table = sheet.iter_rows(values_only=True)
        try:
            header = next(table)
        except StopIteration:
            return []
        names = [_cell_text(name).casefold() for name in header]
        rows = []
        for values in table:
            row = {}
            for index, name in enumerate(names):
                if not name or index >= len(values):
                    continue
                row[name] = values[index]
            rows.append(row)
        return rows
    except Exception:
        return None
    finally:
        workbook.close()


def _combine_mark_rows(csv_rows, workbook_rows):
    """Prefer keep and notes from the workbook, matched back to the file path in the CSV."""
    if workbook_rows is None:
        return _marks_from_rows(csv_rows)
    if any(_cell_text(row.get("file_path")) for row in workbook_rows):
        return _marks_from_rows(workbook_rows)
    from collections import defaultdict

    waiting = defaultdict(list)
    for row in workbook_rows:
        waiting[_loose_key(row)].append(
            (_cell_text(row.get("keep")), _cell_text(row.get("notes")))
        )
    marks = {}
    for row in csv_rows:
        keep = _cell_text(row.get("keep"))
        notes = _cell_text(row.get("notes"))
        queued = waiting.get(_loose_key(row))
        if queued:
            keep, notes = queued.pop(0)
        if keep or notes:
            marks[_mark_key(row)] = (keep, notes)
    return marks


def marks_in_run(folder):
    """Keep and notes saved in one finished run, keyed the same way as a match row."""
    try:
        csv_rows = _read_csv_rows(folder / "matches.csv")
    except (OSError, csv.Error):
        csv_rows = []
    workbook_rows = None
    workbook = folder / "results.xlsx"
    if workbook.is_file():
        workbook_rows = _read_workbook_rows(workbook)
    return _combine_mark_rows(csv_rows, workbook_rows)


def _add_category(names, seen, value):
    text = _cell_text(value)
    if text and text.casefold() not in seen:
        seen.add(text.casefold())
        names.append(text)


def categories_in_run(folder):
    """Categories that run searched, whether or not they had a match."""
    names = []
    seen = set()
    terms_path = folder / "terms_used.csv"
    if terms_path.is_file():
        try:
            with terms_path.open(newline="", encoding="utf-8-sig") as handle:
                for row in csv.DictReader(handle):
                    _add_category(names, seen, row.get("category"))
        except (OSError, csv.Error):
            names, seen = [], set()
        if names:
            return names
    try:
        for row in _read_csv_rows(folder / "matches.csv"):
            _add_category(names, seen, row.get("category"))
    except (OSError, csv.Error):
        pass
    if names:
        return names
    workbook = folder / "results.xlsx"
    if workbook.is_file():
        for row in _read_workbook_rows(workbook) or []:
            _add_category(names, seen, row.get("category"))
    if names:
        return names
    settings = folder / "run_settings.txt"
    if settings.is_file():
        try:
            for line in settings.read_text(encoding="utf-8").splitlines():
                if line.startswith("Category: "):
                    category = line[len("Category: ") :].strip()
                    if category and category.casefold() != "all categories":
                        _add_category(names, seen, category)
                    break
        except OSError:
            pass
    return names


def marks_by_category(results_folder, destination, categories):
    """For each category, marks from the newest earlier run that included it.

    Returns category -> (earlier folder name or None, mark map).
    """
    runs = earlier_runs(results_folder, destination)
    included = {}
    saved = {}
    chosen = {}
    for category in categories:
        source_name = None
        marks = {}
        folded = category.casefold()
        for run in runs:
            if run not in included:
                included[run] = {name.casefold() for name in categories_in_run(run)}
            if folded not in included[run]:
                continue
            if run not in saved:
                saved[run] = marks_in_run(run)
            source_name = run.name
            marks = saved[run]
            break
        chosen[category] = (source_name, marks)
    return chosen


def _match_word(count):
    return "1 match" if count == 1 else f"{count} matches"


def carry_pairs(categories, chosen, counts):
    """One setting per category: where its marks came from, and how many."""
    pairs = []
    for category in categories:
        source_name, _marks = chosen.get(category, (None, {}))
        count = counts.get(category, 0)
        if source_name:
            value = f"carried forward from {source_name}, {_match_word(count)}"
        else:
            value = f"no earlier run included this category, {_match_word(count)}"
        pairs.append((f"Hand-coding marks for {category}", value))
    return pairs


def keep_class(value):
    """yes/y kept, no/n not kept, blank not yet coded, anything else unclear."""
    text = _cell_text(value).casefold()
    if text in {"yes", "y"}:
        return "kept"
    if text in {"no", "n"}:
        return "not_kept"
    if not text:
        return "uncoded"
    return "unclear"


def _keep_tally(matches):
    tally = {"kept": 0, "not_kept": 0, "uncoded": 0, "unclear": 0}
    for match in matches:
        tally[keep_class(match.get("keep"))] += 1
    return tally


def _write_rows(path, fieldnames, rows):
    with Path(path).open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


_MATCH_WIDTHS = {
    "keep": 16,
    "notes": 32,
    "sentence": 80,
}
_NARROW_WIDTH = 14


def _excel_value(value):
    if value is None:
        return ""
    return value


def _wrapped_lines(text, width):
    width = max(int(width) - 1, 8)
    parts = str(text or "").splitlines() or [""]
    total = 0
    for part in parts:
        total += max(1, (len(part) + width - 1) // width)
    return total


def _format_sheet(
    sheet,
    column_count,
    row_count,
    widths=None,
    wrap_names=None,
    header=None,
    default_width=None,
    fit_wrapped=False,
):
    """Bold frozen header, filters, and optional column widths."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    last_row = max(row_count, 1)
    last_column = get_column_letter(max(column_count, 1))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    wrap_names = set(wrap_names or [])
    widths = widths or {}
    wrap_at = {}
    for index, name in enumerate(header or [], start=1):
        letter = get_column_letter(index)
        if name in widths:
            column_width = widths[name]
        elif default_width is not None:
            column_width = default_width
        else:
            column_width = min(max(len(str(name)) + 2, 14), 36)
        sheet.column_dimensions[letter].width = column_width
        if name in wrap_names:
            wrap_at[index] = column_width
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    if fit_wrapped and wrap_at:
        for row_index in range(2, sheet.max_row + 1):
            lines = 1
            for index, column_width in wrap_at.items():
                lines = max(
                    lines,
                    _wrapped_lines(sheet.cell(row_index, index).value, column_width),
                )
            sheet.row_dimensions[row_index].height = 15 * lines


def _fill_sheet(
    workbook,
    title,
    fieldnames,
    rows,
    widths=None,
    wrap_names=None,
    default_width=None,
    fit_wrapped=False,
):
    sheet = workbook.create_sheet(title)
    sheet.append(list(fieldnames))
    for row in rows:
        sheet.append([_excel_value(row.get(name, "")) for name in fieldnames])
    _format_sheet(
        sheet,
        len(fieldnames),
        sheet.max_row,
        widths=widths,
        wrap_names=wrap_names,
        header=fieldnames,
        default_width=default_width,
        fit_wrapped=fit_wrapped,
    )
    return sheet


def _special_column(columns, exact, contains=""):
    """A metadata column with this name, or one whose name contains the word."""
    for column in columns:
        if column.casefold() == exact:
            return column
    if contains:
        for column in columns:
            if contains in column.casefold():
                return column
    return ""


def _placed_name(column, path_column):
    if not column or column == path_column:
        return ""
    return output_column_name(column)


def match_sheet_fields(match_fields, group_name, date_name, title_name):
    """Reading order for the Matches sheet. The file path stays in the CSV only."""
    available = [name for name in match_fields if name != "file_path"]
    present = set(available)
    front = []

    def add(name):
        if name and name in present and name not in front:
            front.append(name)

    add(group_name)
    add(date_name)
    add(title_name)
    for name in (
        "term",
        "sentence",
        "keep",
        "notes",
        "category",
        "language",
        "short",
        "location",
        "file_name",
    ):
        add(name)
    for name in available:
        add(name)
    body = [name for name in front if name.casefold() not in {"discussion", "url"}]
    discussion = [name for name in front if name.casefold() == "discussion"]
    url = [name for name in front if name.casefold() == "url"]
    return body + discussion + url


def _write_workbook(
    path,
    match_rows,
    sheet_fields,
    category_rows,
    group_rows,
    term_rows,
    decision_rows,
    context_table,
    settings_pairs,
):
    """One workbook with one sheet per table. CSV files are written separately."""
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "Excel workbooks need the openpyxl library. Install the packages in requirements.txt."
        ) from error

    category_fields = [
        "category",
        "matches",
        "documents",
        "groups",
        "matches_kept",
        "matches_not_kept",
        "matches_uncoded",
        "matches_unclear",
    ]
    group_fields = [
        "category",
        "group",
        "matches",
        "words",
        "matches_per_10000_words",
        "matches_kept",
        "matches_not_kept",
        "matches_uncoded",
        "matches_unclear",
        "matches_kept_per_10000_words",
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    _fill_sheet(
        workbook,
        "Matches",
        sheet_fields,
        match_rows,
        widths=_MATCH_WIDTHS,
        wrap_names=["sentence"],
        default_width=_NARROW_WIDTH,
        fit_wrapped=True,
    )
    _fill_sheet(workbook, "Counts by category", category_fields, category_rows)
    if group_rows is not None:
        _fill_sheet(workbook, "Counts by group", group_fields, group_rows)
    _fill_sheet(workbook, "Terms used", ["category", "term", "stem"], term_rows)
    _fill_sheet(workbook, "Word decisions", list(DECISION_COLUMNS), decision_rows)
    if context_table is not None:
        _fill_sheet(workbook, "Context words", list(CONTEXT_COLUMNS), context_table)
    setting_rows = []
    for item in settings_pairs:
        if item is None:
            continue
        setting, value = item
        setting_rows.append({"Setting": setting, "Value": "" if value is None else value})
    _fill_sheet(
        workbook,
        "Run settings",
        ["Setting", "Value"],
        setting_rows,
        widths={"Setting": 72, "Value": 72},
        wrap_names=["Setting", "Value"],
        fit_wrapped=True,
    )
    workbook.save(path)


def _folder_category(category):
    name = (category or "").strip() or "All categories"
    if name.casefold() == "all categories":
        name = "All categories"
    cleaned = re.sub(r'[<>:"/\\|?*]+', "-", name).strip(" .")
    return cleaned or "category"


def _run_folder(results_root, category, when):
    """Date, minute, and category. Add the seconds only when that name is taken."""
    label = _folder_category(category)
    minute = when.strftime("%Y-%m-%d_%H-%M")
    destination = results_root / f"{minute}_{label}"
    if not destination.exists():
        return destination
    second = when.strftime("%Y-%m-%d_%H-%M-%S")
    destination = results_root / f"{second}_{label}"
    suffix = 2
    while destination.exists():
        destination = results_root / f"{second}_{label}-{suffix}"
        suffix += 1
    return destination


def _remove_unfinished_run(destination, results_root):
    """Delete one run folder that was still being written. Leave every other folder."""
    try:
        folder = destination.resolve()
        root = results_root.resolve()
    except OSError:
        return
    if folder == root or root not in folder.parents or not folder.is_dir():
        return
    shutil.rmtree(folder)


def write_run(
    config,
    corpus,
    matches,
    categories,
    terms,
    decision_categories=None,
    context_rows=None,
    should_stop=None,
):
    """Create a new dated folder and return its path. Older runs are left as they were."""
    results_root = Path(config.results_folder)
    results_root.mkdir(parents=True, exist_ok=True)
    check_stop(should_stop)
    destination = _run_folder(results_root, config.category, datetime.now())
    destination.mkdir()
    try:
        return _write_run_files(
            config,
            corpus,
            matches,
            categories,
            terms,
            destination,
            results_root,
            decision_categories,
            context_rows,
            should_stop,
        )
    except RunStopped:
        _remove_unfinished_run(destination, results_root)
        raise


def _write_run_files(
    config,
    corpus,
    matches,
    categories,
    terms,
    destination,
    results_root,
    decision_categories,
    context_rows,
    should_stop,
):
    check_stop(should_stop)
    chosen = marks_by_category(results_root, destination, categories)
    path_column = (config.path_column or "").strip()
    meta_columns = [
        column
        for column in metadata_columns(corpus.documents)
        if column != path_column
    ]
    renamed_meta = [(column, output_column_name(column)) for column in meta_columns]
    match_fields = (
        ["file_name", "file_path"]
        + [renamed for _original, renamed in renamed_meta]
        + ["language", "short", "category", "term", "location", "sentence", "keep", "notes"]
    )
    match_rows = []
    carried = {category: 0 for category in categories}
    for match in matches:
        row = {
            "file_name": match["file_name"],
            "file_path": match["file_path"],
            "language": match["language"],
            "short": "yes" if match["short"] else "no",
            "category": match["category"],
            "term": match["term"],
            "location": match["location"],
            "sentence": match["sentence"],
            "keep": "",
            "notes": "",
        }
        for original, renamed in renamed_meta:
            row[renamed] = match["metadata"].get(original, "")
        _source, marks = chosen.get(match["category"], (None, {}))
        keep, notes = marks.get(_mark_key(row), ("", ""))
        if keep or notes:
            row["keep"] = keep
            row["notes"] = notes
            if match["category"] in carried:
                carried[match["category"]] += 1
        match["keep"] = row["keep"]
        match_rows.append(row)
    _write_rows(destination / "matches.csv", match_fields, match_rows)
    check_stop(should_stop)

    counted_matches = [
        match for match in matches if include_in_counts(match["language"], config.count_language)
    ]
    counted_documents = [
        document
        for document in corpus.documents
        if include_in_counts(document.language, config.count_language)
    ]
    grouping = (config.group_column or "").strip()
    category_fields = [
        "category",
        "matches",
        "documents",
        "groups",
        "matches_kept",
        "matches_not_kept",
        "matches_uncoded",
        "matches_unclear",
    ]
    category_rows = category_count_rows(categories, counted_matches, grouping)
    _write_rows(destination / "counts_by_category.csv", category_fields, category_rows)
    group_fields = [
        "category",
        "group",
        "matches",
        "words",
        "matches_per_10000_words",
        "matches_kept",
        "matches_not_kept",
        "matches_uncoded",
        "matches_unclear",
        "matches_kept_per_10000_words",
    ]
    group_rows = []
    if grouping:
        group_rows = group_count_rows(
            categories, counted_matches, counted_documents, grouping
        )
        _write_rows(destination / "counts_by_group.csv", group_fields, group_rows)
    check_stop(should_stop)

    write_terms(destination / "terms_used.csv", terms)
    names = categories if decision_categories is None else decision_categories
    wanted = {category.casefold() for category in names}
    decision_rows = [
        row
        for row in read_decision_rows(config.term_list)
        if (row.get("category") or "").casefold() in wanted
    ]
    _write_rows(destination / "word_decisions.csv", list(DECISION_COLUMNS), decision_rows)
    check_stop(should_stop)
    if getattr(config, "context_words", True):
        _write_rows(
            destination / "context_words.csv",
            CONTEXT_COLUMNS,
            context_rows or [],
        )
        context_note = CONTEXT_NOTE
    else:
        context_note = CONTEXT_SKIPPED_NOTE
    carried_lines = carry_pairs(categories, chosen, carried)
    settings_text, settings_pairs = run_record(
        config,
        corpus,
        destination,
        len(matches),
        len(counted_matches),
        context_note,
        carried_lines,
    )
    (destination / "run_settings.txt").write_text(settings_text, encoding="utf-8")
    check_stop(should_stop)
    term_rows = [
        {"category": term.category, "term": term.term, "stem": term.stem} for term in terms
    ]
    original_meta = [original for original, _renamed in renamed_meta]
    date_name = _placed_name(
        _special_column(original_meta, "date", contains="date"), path_column
    )
    title_name = _placed_name(_special_column(original_meta, "title"), path_column)
    group_name = _placed_name(grouping, path_column)
    _write_workbook(
        destination / "results.xlsx",
        match_rows,
        match_sheet_fields(match_fields, group_name, date_name, title_name),
        category_rows,
        group_rows if grouping else None,
        term_rows,
        decision_rows,
        (context_rows or []) if getattr(config, "context_words", True) else None,
        settings_pairs,
    )
    return destination


def category_count_rows(categories, matches, grouping):
    rows = []
    for category in categories:
        chosen = [match for match in matches if match["category"] == category]
        documents = {match["file_path"] for match in chosen}
        if grouping:
            groups = {group_value_from_match(match, grouping) for match in chosen}
            group_count = len(groups)
        else:
            group_count = ""
        tally = _keep_tally(chosen)
        rows.append(
            {
                "category": category,
                "matches": len(chosen),
                "documents": len(documents),
                "groups": group_count,
                "matches_kept": tally["kept"],
                "matches_not_kept": tally["not_kept"],
                "matches_uncoded": tally["uncoded"],
                "matches_unclear": tally["unclear"],
            }
        )
    return rows


def group_value_from_match(match, column):
    value = (match["metadata"].get(column) or "").strip()
    return value or "(blank)"


def group_count_rows(categories, matches, documents, grouping):
    groups = []
    for document in documents:
        value = group_value(document, grouping)
        if value not in groups:
            groups.append(value)
    words_by_group = {value: 0 for value in groups}
    for document in documents:
        words_by_group[group_value(document, grouping)] += document.word_count

    rows = []
    for category in categories:
        chosen = [match for match in matches if match["category"] == category]
        for value in groups:
            count = sum(1 for match in chosen if group_value_from_match(match, grouping) == value)
            words = words_by_group[value]
            grouped = [
                match for match in chosen if group_value_from_match(match, grouping) == value
            ]
            tally = _keep_tally(grouped)
            rows.append(
                {
                    "category": category,
                    "group": value,
                    "matches": count,
                    "words": words,
                    "matches_per_10000_words": per_10000(count, words),
                    "matches_kept": tally["kept"],
                    "matches_not_kept": tally["not_kept"],
                    "matches_uncoded": tally["uncoded"],
                    "matches_unclear": tally["unclear"],
                    "matches_kept_per_10000_words": per_10000(tally["kept"], words),
                }
            )
    return rows


def _bullet_list(items, empty_note):
    if not items:
        return [empty_note]
    lines = []
    for item in items:
        if isinstance(item, tuple):
            lines.append(f"- {item[0]} ({item[1]})")
        else:
            lines.append(f"- {item}")
    return lines


def _settings_text(pairs):
    lines = []
    for item in pairs:
        if item is None:
            lines.append("")
            continue
        setting, value = item
        if value is None:
            lines.append(setting)
        else:
            lines.append(f"{setting}: {value}")
    return "\n".join(lines)


def run_record(
    config, corpus, destination, match_count, counted_match_count, context_note, carried_lines
):
    header_fields = ", ".join(config.header_fields) if config.header_fields else "none"
    grouping = config.group_column or "none"
    pairs = [
        ("Results folder", str(destination)),
        ("Corpus folder", config.corpus_folder),
        ("Term list", config.term_list),
        ("Word decisions file", str(decisions_path(config.term_list))),
        ("Metadata file", config.metadata_file or "none"),
        ("File-path column", config.path_column or "none"),
        ("Group column", grouping),
        ("Category", config.category),
        (
            "Remove formatting and web addresses",
            "yes" if config.strip_formatting else "no",
        ),
        ("Read opening header lines", "yes" if config.use_headers else "no"),
        (
            "Header fields searched",
            header_fields if config.use_headers else "not used",
        ),
        (
            "Flag documents shorter than",
            f"{int(config.short_word_limit)} words",
        ),
        ("Language used for the counts", config.count_language),
        None,
        ("The sentence file, matches.csv, includes every language.", None),
        ("The count files include only the language chosen above.", None),
        ("All languages means every document is included in the counts.", None),
        (context_note, None),
        *carried_lines,
        ("Matches per 10,000 words uses the words in the text that was searched.", None),
        ("A blank rate means that group had no words to divide by.", None),
        ("The groups column is blank when no grouping column was chosen.", None),
        None,
        ("Documents read", str(len(corpus.documents))),
        (
            "Documents flagged as short",
            str(sum(1 for document in corpus.documents if document.short)),
        ),
        ("Matches in the sentence file", str(match_count)),
        ("Matches included in the counts", str(counted_match_count)),
        None,
        ("Files listed in the metadata that could not be found:", None),
        *[(line, None) for line in _bullet_list(corpus.missing, "- none")],
        None,
        ("Files that could not be read:", None),
        *[(line, None) for line in _bullet_list(corpus.unreadable, "- none")],
        None,
        ("Files with no text, skipped:", None),
        *[(line, None) for line in _bullet_list(corpus.empty, "- none")],
        None,
    ]
    return _settings_text(pairs), pairs
